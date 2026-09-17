# -*- coding: utf-8 -*-
"""Teste de fumaça do pipeline — sem dependência de pytest.

    python testes/test_pipeline.py

Cobre o que já quebrou de verdade neste projeto, para não quebrar de novo:

 1. todos os módulos importam;
 2. o validador de slug aceita o padrão e recusa acento, espaço, maiúscula e data inválida;
 3. a medição do bruto acha o corpo (a linha mais longa) e não chuta o cabeçalho;
 4. o montador DOCX preserva parágrafos, negrito e nota em itálico;
 5. o conversor Markdown não deixa asterisco à mostra em ênfase aninhada
    (**negrito com `código` dentro** e ***negrito-itálico*** já saíram errados);
 6. formas proibidas são varridas COM fronteira de palavra — "Demiurg" não pode
    acusar dentro de "Demiurgo", nem "enoteísmo" dentro de "henoteísmo";
 7. os portões rápidos do QA continuam verdes na transcrição de referência;
 8. o catálogo mede a transcrição de referência como ela é;
 9. o modelo de pasta continua íntegro (é dele que toda transcrição nova nasce);
10. a curadoria da KB é cirúrgica: atesta antes de gravar, não duplica, não estraga o
    markdown da ficha e respeita o status da fila (aplicada não se reaplica);
11. o padrão Y fecha o ciclo da fonte audiovisual: termo novo nasce citando fonte
    registrada, fronteira de palavra serve a "/Kaggen", e os três lugares que guardam o
    URL (biblio.json, metadados.yaml, midia/README.md) concordam entre si;
12. o perfil de motor STT (`rc_perfil_stt.py`) mede o que diz medir: marcador oral não
    confunde "ó" com o artigo "o", equivalência conceitual não vira corrupção, e um
    arquivo paragraphado dispara as quebras de suposição da esteira em vez de passar por
    compatível em silêncio;
13. proveniência e ruído: o instrumento sabe dizer se dois arquivos vêm do mesmo
    reconhecimento de fala, e separa inicial de frase de candidata real a termo novo;
14. arquitetura de três camadas: o corpo do bruto tem critério único e avisado (rc_leitura),
    e o portão G9 falha em cada um dos cinco modos de o derivado estragar o trabalho —
    sha256 trocado, divergência acima do teto, censura não restaurada, asterisco no texto
    revisado e máscara sem registro.
"""
from __future__ import annotations

import hashlib
import re
import shutil
import sys
import tempfile
from collections import Counter
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "ferramentas"))

FIXTURE = RAIZ / "testes" / "fixtures" / "mini-transcricao.txt"
REFERENCIA = RAIZ / "transcricoes" / "2026-09-14-revelacoes-cosmicas-urgente"

_falhas: list[str] = []
_ok = 0


def verificar(nome: str, condicao: bool, detalhe: str = "") -> None:
    global _ok
    if condicao:
        _ok += 1
        print(f"  ok   {nome}")
    else:
        _falhas.append(f"{nome} — {detalhe}")
        print(f"  FALHA {nome} — {detalhe}")


# --------------------------------------------------------------------------------------
print("\n1. módulos importam")
modulos = {}
for nome in ["rc_kb", "rc_lexicon", "rc_variantes", "rc_diagnostico", "rc_docx",
             "md_para_docx", "rc_qa", "rc_indice", "rc_novo"]:
    try:
        modulos[nome] = __import__(nome)
        verificar(f"import {nome}", True)
    except Exception as exc:  # noqa: BLE001
        verificar(f"import {nome}", False, str(exc))

DX = modulos.get("rc_docx")
MD = modulos.get("md_para_docx")
QA = modulos.get("rc_qa")
IND = modulos.get("rc_indice")
NOVO = modulos.get("rc_novo")

# --------------------------------------------------------------------------------------
print("\n2. validador de slug")
if NOVO:
    bons = ["2026-10-02-lemuria-terry-fabris", "2027-01-15-jan-val-ellam", "2026-12-31-a"]
    for s in bons:
        verificar(f"aceita '{s}'", NOVO.validar_slug(s) is None, str(NOVO.validar_slug(s)))
    ruins = {
        "Revelações Cósmicas": "acento e maiúscula",
        "2026-10-02 titulo": "espaço",
        "2026-10-02-Titulo": "maiúscula",
        "lemuria-terry-fabris": "sem data",
        "2026-02-30-lemuria": "data inexistente",
        "2026-10-02-" + "x" * 70: "longo demais",
    }
    for s, por_que in ruins.items():
        verificar(f"recusa '{s[:34]}' ({por_que})", NOVO.validar_slug(s) is not None, "foi aceito")

# --------------------------------------------------------------------------------------
print("\n3. medição do bruto")
if NOVO:
    m = NOVO.medir_bruto(FIXTURE)
    verificar("sha256 tem 64 hex", re.fullmatch(r"[0-9a-f]{64}", m["sha256"]) is not None, m["sha256"])
    verificar("corpo detectado na linha 13", m["corpo_linha"] == 13, f"achou {m['corpo_linha']}")
    verificar("cabeçalho com 12 linhas", m["linhas_cabecalho"] == 12, f"achou {m['linhas_cabecalho']}")
    verificar("corpo tem palavras", m["corpo_palavras"] > 50, str(m["corpo_palavras"]))
    verificar("bytes batem com o arquivo", m["bytes"] == FIXTURE.stat().st_size)

# --------------------------------------------------------------------------------------
print("\n4. montador DOCX")
if DX:
    from docx import Document
    with tempfile.TemporaryDirectory() as tmp:
        md = Path(tmp) / "b.md"
        md.write_text("## Bloco 1 — teste\n\n**[JAN VAL ELLAM]** Texto com **negrito** "
                      "e uma nota [NOTA: no bruto, \"Xavé\"].\n", encoding="utf-8")
        itens = DX.analisar_blocos([md], [])
        saida = DX.montar(itens, Path(tmp) / "s.docx", "Título", "Subtítulo")
        verificar("arquivo gerado", saida.exists())
        d = Document(str(saida))
        textos = [p.text for p in d.paragraphs]
        verificar("heading presente", any(t.startswith("Bloco 1") for t in textos), str(textos[:4]))
        corpo = next(p for p in d.paragraphs if "negrito" in p.text)
        verificar("negrito aplicado", any(r.bold and r.text == "negrito" for r in corpo.runs))
        verificar("nota em itálico", any(r.italic and r.text.startswith("[NOTA") for r in corpo.runs))
        # o validador não pode punir a nota que cita o bruto
        verificar("nota não conta como sobrevivência",
                  DX.MARCADOR_RE.sub(" ", corpo.text).count("Xavé") == 0)

# --------------------------------------------------------------------------------------
print("\n5. conversor Markdown — ênfase aninhada")
if MD:
    from docx import Document
    with tempfile.TemporaryDirectory() as tmp:
        md = Path(tmp) / "d.md"
        md.write_text("# Título\n\n**Por que `bloco-01.md` e não `bloco-1.md`:** o glob "
                      "`bloco-*.md` ordena errado.\n\n- ***A Divina Colmeia*** e *Deep Utopia*.\n\n"
                      "```\n*.docx binary\n```\n", encoding="utf-8")
        MD.converter(md, Path(tmp) / "d.docx", "T", "S")
        d = Document(str(Path(tmp) / "d.docx"))
        for p in d.paragraphs:
            em_code = any(r.font.name == "Consolas" for r in p.runs)
            if em_code:
                continue
            verificar(f"sem markup à mostra em '{p.text[:38]}…'",
                      "*" not in p.text and "`" not in p.text, p.text[:80])
        negrito_codigo = next((p for p in d.paragraphs if "Por que" in p.text), None)
        verificar("negrito com código dentro",
                  bool(negrito_codigo) and any(r.bold and r.font.name == "Consolas" for r in negrito_codigo.runs))
        bi = next((p for p in d.paragraphs if "Divina Colmeia" in p.text), None)
        verificar("negrito-itálico (***…***)",
                  bool(bi) and any(r.bold and r.italic for r in bi.runs))

# --------------------------------------------------------------------------------------
print("\n6. fronteira de palavra nas formas proibidas")
if QA:
    # o texto contém só as formas corretas: as truncadas aparecem aqui apenas como
    # padrão de busca — se casarem, é porque a varredura está sem fronteira de palavra
    texto = "O Demiurgo e o henoteísmo aparecem aqui; as formas truncadas não devem casar."
    sem_marcador = QA.MARCADOR_RE.sub(" ", texto)
    for forma in ["Demiurg", "enoteísmo"]:
        n = len(re.findall(r"\b" + re.escape(forma) + r"\b", sem_marcador, re.I))
        verificar(f"'{forma}' não casa dentro da palavra correta", n == 0, f"{n} ocorrências")
    for forma in ["Demiurgo", "henoteísmo"]:
        n = len(re.findall(r"\b" + re.escape(forma) + r"\b", sem_marcador, re.I))
        verificar(f"'{forma}' casa", n == 1, f"{n} ocorrências")
    proibidas = QA.formas_proibidas(REFERENCIA, RAIZ / "KB-RC")
    verificar("formas proibidas derivadas da KB + ledger", len(proibidas) >= 10, f"{len(proibidas)}")
    verificar("Quarentena 'Sofia' capturada", "Sofia" in proibidas, str(sorted(proibidas)[:5]))
    # quase-canônicas do corretor do editor (vídeo 2, fila 0049/0050, despacho de 17/09/2026)
    verificar("Quarentena 'glues' capturada (RC-034)", "glues" in proibidas, str(sorted(proibidas)[:6]))
    verificar("Quarentena 'pósetron' capturada (RC-636)", "pósetron" in proibidas)

    # uma [NOTA] pode citar a notação da KB, que tem colchete dentro: «pósitrons [STT 'positelétron']».
    # Sem tolerar o aninhamento o marcador fecha cedo, o rabo volta a ser corpo e a forma proibida
    # citada como evidência vira falso positivo no G3 (defeito achado no vídeo 2, bloco 3).
    aninhado = 'corpo [NOTA: o editor trocou por "pósetron"; RC-636 registra "pósitrons [STT x]", fim] mais corpo'
    expurgo = QA.MARCADOR_RE.sub(" ", aninhado)
    verificar("marcador com colchete aninhado é expurgado por inteiro",
              "pósetron" not in expurgo and "[" not in expurgo, expurgo)
    n = DX.NOTA_RE.search("[NOTA: a [STT b] c]")
    verificar("NOTA_RE casa a nota inteira quando há colchete dentro",
              bool(n) and n.group(0) == "[NOTA: a [STT b] c]", n.group(0) if n else "sem casamento")

# --------------------------------------------------------------------------------------
print("\n7. portões rápidos na transcrição de referência")
if QA and REFERENCIA.exists():
    meta = QA.ler_metadados(REFERENCIA)
    verificar("metadados.yaml lido", meta.get("slug") == REFERENCIA.name, str(meta.get("slug")))
    for codigo, nome, funcao in QA.PORTOES:
        if codigo not in {"G1", "G2", "G4", "G7"}:
            continue
        status, detalhe = funcao(REFERENCIA, meta)
        verificar(f"{codigo} {nome}", status == QA.OK, f"{status}: {detalhe}")
    verificar("estágio medido no disco", QA.estagio(REFERENCIA) == "40-devolvida",
              QA.estagio(REFERENCIA))

# --------------------------------------------------------------------------------------
print("\n8. catálogo mede a referência")
if IND and REFERENCIA.exists():
    linha = IND.medir(REFERENCIA)
    verificar("8 blocos", linha["blocos"] == 8, str(linha["blocos"]))
    verificar("18.966 palavras revisadas", linha["palavras_revisadas"] == 18966, str(linha["palavras_revisadas"]))
    verificar("32 notas", linha["notas"] == 32, str(linha["notas"]))
    verificar("status 40-devolvida", linha["estatus"] == "40-devolvida", linha["estatus"])
    verificar("falantes listados", "JAN VAL ELLAM" in linha["falantes"], linha["falantes"])
    verificar("artefatos do estágio presentes",
              IND.artefatos_faltando(REFERENCIA, linha["estatus"]) == [])

# --------------------------------------------------------------------------------------
print("\n9. modelo de pasta íntegro")
MODELO = RAIZ / "transcricoes" / "_modelo"
esperados = ["README.md", "00-fonte/metadados.yaml", "00-fonte/midia/README.md",
             "10-diagnostico/README.md", "20-blocos/notas-de-revisao.md",
             "30-produto/README.md", "40-devolucao/README.md",
             "90-registro/diario-de-bordo.md", "90-registro/despachos/README.md"]
for rel in esperados:
    verificar(f"_modelo/{rel}", (MODELO / rel).exists())
if (MODELO / "00-fonte" / "metadados.yaml").exists():
    modelo_txt = (MODELO / "00-fonte" / "metadados.yaml").read_text(encoding="utf-8")
    verificar("modelo tem marcadores {{…}} para substituir", "{{slug}}" in modelo_txt)

# --------------------------------------------------------------------------------------
print("\n10. curadoria da KB (rc_curadoria)")
import rc_curadoria as CU  # noqa: E402
import rc_kb as KB10  # noqa: E402
import rc_lexicon as L10  # noqa: E402

# a atestação tem fronteira de palavra: "Demiurg" não pode casar dentro de "Demiurgo"
verificar("atestar usa fronteira de palavra",
          CU.atestar(L10.norm("o Demiurgo e também Demiurg"), "Demiurg") == 1)

# o separador de seção é sempre uma linha em branco — nem zero, nem duas
_base = "# Título\n\nTexto.\n## Outra\n"
_res = CU._inserir(_base, len("# Título\n\nTexto.\n"), "## Nova\n- item")
verificar("_inserir separa com exatamente uma linha em branco",
          "\n\n## Nova\n- item\n\n## Outra\n" in _res and "\n\n\n" not in _res)

# a fila: item aplicado não volta, item bloqueado só entra se pedido
_linhas = [{"id": "1", "tipo": "nova-variante", "status": "aplicada"},
           {"id": "2", "tipo": "nova-variante", "status": "pendente"},
           {"id": "3", "tipo": "novo-termo", "status": "pendente"},
           {"id": "4", "tipo": "nova-variante", "status": "bloqueada"}]
verificar("selecionar ignora aplicada e bloqueada",
          [l["id"] for l in CU.selecionar(_linhas, "nova-variante", [])] == ["2"])
verificar("selecionar inclui bloqueada quando pedido",
          [l["id"] for l in CU.selecionar(_linhas, "nova-variante", [], True)] == ["2", "4"])
verificar("selecionar por id ignora status (revisão pontual)",
          [l["id"] for l in CU.selecionar(_linhas, "nova-variante", ["1"])] == ["1"])
verificar("variantes_do_item separa por barra",
          CU.variantes_do_item({"termo": "arcontos / erontes"}) == ["arcontos", "erontes"])

# cirurgia numa ficha sintética: cria a seção, registra a proveniência, não estraga nada
_tmp = Path(tempfile.mkdtemp(prefix="rc-curadoria-"))
(_tmp / "termos").mkdir()
_ficha = _tmp / "termos" / "RC-999-termo-de-teste.md"
_ficha.write_text(
    '+++\ncodigo = "RC-999"\nnome = "Termo de Teste"\nstatus = "candidato"\n'
    'atualizado = "2020-01-01"\n+++\n'
    "# Termo de Teste — RC-999\n\n## Definição Sintética\nDefinição.\n\n"
    "## Ampliação\n### P2020-01-01\n- citação.\n", encoding="utf-8")
_f = KB10.carregar_ficha(_ficha)
_r1 = CU.aplicar_variantes(_ficha, _f, ["Termo Deteçte", "termo de teste"],
                           "### nota\n- evidência da transcrição de teste.")
_txt = _ficha.read_text(encoding="utf-8")
verificar("variante nova gravada", "Termo Deteçte" in _r1["adicionadas"])
verificar("variante igual ao canônico é recusada",
          any("canônica" in m for _, m in _r1["puladas"]))
verificar("seção Etimologia criada quando a ficha não a tem", _r1["secao_criada"])
_f2 = KB10.carregar_ficha(_ficha)
verificar("parser da KB enxerga a variante gravada", "Termo Deteçte" in _f2.variacoes_stt)
verificar("proveniência registrada em Atualização",
          "evidência da transcrição de teste" in _f2.secoes.get("Atualização", ""))
verificar("frontmatter atualizado sai do valor velho", 'atualizado = "2020-01-01"' not in _txt)
verificar("sem cabeçalho colado na linha anterior", not re.search(r"[^\n]\n##\s", _txt))
verificar("sem três linhas em branco seguidas", "\n\n\n" not in _txt)

# idempotência: rodar de novo não acrescenta nada nem toca no arquivo
_r2 = CU.aplicar_variantes(_ficha, KB10.carregar_ficha(_ficha), ["Termo Deteçte"], "### nota\n- x")
verificar("reaplicar não duplica variante", _r2["adicionadas"] == [])
verificar("reaplicar não modifica o arquivo",
          _ficha.read_text(encoding="utf-8") == _txt)

# a fila real: o que está marcado como aplicado tem data e curador
_fila, _ = CU.carregar_fila(CU.FILA_PADRAO)
_aplicadas = [l for l in _fila if l["status"] == "aplicada"]
verificar("fila tem itens aplicados com data e curador",
          bool(_aplicadas) and all(l["data_aplicada"] and l["curador"] for l in _aplicadas))
_i23 = next((l for l in _fila if l["id"] == "0023"), None)
verificar("item desbloqueado registra o motivo na própria evidencia",
          _i23 is not None and _i23["status"] == "aplicada"
          and "DESBLOQUEADA" in _i23["evidencia"] and _i23["codigo_afetado"] == "RC-952")
verificar("nenhum item aplicado sem variante legível",
          all(CU.variantes_do_item(l) for l in _aplicadas))

# --------------------------------------------------------------------------------------
print("\n11. padrão Y (fonte audiovisual) e criação de termos")
import json as J11  # noqa: E402
import rc_lexicon as LX  # noqa: E402
import rc_termo as RT  # noqa: E402
import rc_indice as RI  # noqa: E402

# fronteira de palavra precisa servir a forma que começa em barra — "/Kaggen" é RC-948
verificar("fronteira casa '/Kaggen' isolado",
          len(re.findall(LX.fronteira("/Kaggen"), "chamavam Javé de /Kaggen, que quer dizer")) == 1)
verificar("fronteira não casa dentro de derivada",
          len(re.findall(LX.fronteira("/Kaggen"), "kaggeniano não existe")) == 0)
verificar("fronteira mantém a disciplina Demiurg/henoteísmo",
          len(re.findall(LX.fronteira("Demiurg"), "o Demiurgo e Demiurg")) == 1
          and len(re.findall(LX.fronteira("enoteísmo"), "henoteísmo")) == 0)

# a ferramenta de criar termo valida antes de gravar
_spec = {"meta": {}, "termos": [{"nome": "Termo Falso", "categoria": "Conceitos Cosmológicos",
                                 "subcategoria": "2.1 Cosmogonia / Criação", "status": "provisório",
                                 "confianca_fonte": "média", "fontes": ["Y1999-01-01"],
                                 "definicao": "x", "relacionados": [{"codigo": "RC-999"}]}]}
_canon = {"termos": [{"codigo": "RC-001", "nome": "Javé", "categoria": "Conceitos Cosmológicos",
                      "subcategoria": "2.1 Cosmogonia / Criação"}]}
_, _probs = RT.validar(_spec, _canon, [], {"RC-001": 1})
verificar("validar recusa fonte fora de biblio.json", any("biblio.json" in p for p in _probs))
verificar("validar recusa relacionado inexistente", any("RC-999" in p for p in _probs))
_, _probs2 = RT.validar({"meta": {}, "termos": [dict(_spec["termos"][0], fontes=["Y2026-09-14"],
                                                     relacionados=[])]},
                        _canon, [{"codigo": "Y2026-09-14"}], {"RC-001": 1})
verificar("validar aprova especificação correta", _probs2 == [])
verificar("validar recusa categoria fora da taxonomia",
          any("taxonomia" in p for p in RT.validar(
              {"meta": {}, "termos": [dict(_spec["termos"][0], categoria="Invenções",
                                           fontes=["Y2026-09-14"], relacionados=[])]},
              _canon, [{"codigo": "Y2026-09-14"}], {"RC-001": 1})[1]))
verificar("próximo código livre", RT.proximo_codigo(_canon["termos"]) == "RC-002")
verificar("nome de arquivo ASCII-safe",
          RT.slug_arquivo("RC-948", "/Kaggen (nome san de Javé)") == "RC-948-kaggen-nome-san-de-jave.md"
          and RT.slug_arquivo("RC-955", "Javé 2.0") == "RC-955-jave-2-0.md")

# os dez termos do lote 02 existem e são coerentes
_canon_real = J11.loads((RAIZ / "KB-RC" / "canonico.json").read_text(encoding="utf-8"))
_novos = [t for t in _canon_real["termos"] if t["codigo"] >= "RC-947"]
verificar("lote 02 criou RC-947 a RC-956", len(_novos) == 10)
verificar("todos citam a fonte Y", all(t["fontes"] == ["Y2026-09-14"] for t in _novos))
verificar("nenhum nasceu 'verificado' (fonte STT única)",
          all(t["status"] in {"provisório", "candidato"} for t in _novos))
verificar("todos têm ficha no disco",
          all((RAIZ / "KB-RC" / "termos" / f"RC-{947 + i}").exists() or
              list((RAIZ / "KB-RC" / "termos").glob(f"RC-{947 + i}-*.md")) for i in range(10)))
_obras = J11.loads((RAIZ / "KB-RC" / "biblio.json").read_text(encoding="utf-8"))["obras"]
_y = next((o for o in _obras if o["codigo"] == "Y2026-09-14"), None)
verificar("Y2026-09-14 registrada em biblio.json", _y is not None)
if _y:
    _meta_txt = (REFERENCIA / "00-fonte" / "metadados.yaml").read_text(encoding="utf-8")
    _midia_txt = (REFERENCIA / "00-fonte" / "midia" / "README.md").read_text(encoding="utf-8")
    verificar("os três lugares concordam no URL (biblio, metadados, midia)",
              _y["url"] in _meta_txt and _y["url"] in _midia_txt)
    verificar("registro Y tem duração conferida, não estimada",
              _y.get("duracao_min") == 146 and _y.get("duracao") == "2:25:50")
    verificar("mídia não arquivada", _y.get("midia_arquivada") is False)

# a variante que trocou de ficha no lote 02
_f953 = KB10.carregar_ficha(next((RAIZ / "KB-RC" / "termos").glob("RC-953-*.md")))
_f176 = KB10.carregar_ficha(next((RAIZ / "KB-RC" / "termos").glob("RC-176-*.md")))
verificar("'circuito coméico' está em RC-953", "circuito coméico" in _f953.variacoes_stt)
verificar("'circuito coméico' saiu de RC-176", "circuito coméico" not in _f176.variacoes_stt)
verificar("RC-176 guarda remissiva da mudança", "RC-953" in _f176.secoes.get("Atualização", ""))

# o catálogo fiscaliza o padrão Y: problema quando maduro, aviso quando em captura
_p1, _a1 = RI.checar_fonte_y({"slug": "s", "estatus": "40-devolvida", "url": ""}, REFERENCIA)
verificar("sem URL em estágio maduro é problema", bool(_p1) and not _a1)
_p2, _a2 = RI.checar_fonte_y({"slug": "s", "estatus": "00-fonte", "url": ""}, REFERENCIA)
verificar("sem URL em captura é aviso (não reprova pasta nova)", not _p2 and bool(_a2))
_p3, _a3 = RI.checar_fonte_y({"slug": "2026-09-14-revelacoes-cosmicas-urgente",
                              "estatus": "40-devolvida", "url": _y["url"]}, REFERENCIA)
verificar("URL registrada e apontando de volta passa limpa", not _p3 and not _a3)

# 12. o perfil de motor STT mede o que diz medir — e avisa quando a esteira não cabe no arquivo.
#     Nasceu do experimento de 16/09/2026 (YouTube × NotebookLM): trocar o motor de STT muda a
#     FORMA do arquivo de entrada, e três suposições da esteira dependem dessa forma.
import rc_perfil_stt as PS  # noqa: E402
import rc_lexicon as L11  # noqa: E402

FIX_PARA = RAIZ / "testes" / "fixtures" / "stt-com-paragrafos-sintetico.txt"

# o bug que a régua tinha: "ó" normalizado vira "o" e contava artigo como marcador oral
verificar("'ó' vocativo não é contado como o artigo 'o'",
          PS.contar_marcador("ó", "o gato viu o cão, ó fulano", L11.norm("o gato viu o cão, ó fulano")) == 1)
verificar("'né' é contado com e sem acento (motor que não acentua)",
          PS.contar_marcador("né", "é isso ne, é isso né", L11.norm("é isso ne, é isso né")) == 2)

_txt_yt, _txt_par = PS.ler(FIXTURE)["texto"], PS.ler(FIX_PARA)["texto"]
_perfis_fix = [
    {"assinatura": PS.assinatura(_txt_yt), "texto_bruto": _txt_yt,
     "disfluencia": PS.eixo_disfluencia(_txt_yt)},
    {"assinatura": PS.assinatura(_txt_par), "texto_bruto": _txt_par.replace("Brahma", "B****"),
     "disfluencia": PS.eixo_disfluencia(_txt_par)},
]
_para = PS.ler(FIX_PARA)
_int_para = PS.eixo_integracao(_para)
_pont_para = PS.eixo_pontuacao(_para["texto"], PS.paragrafos(_para["linhas"]))
verificar("fixture paragraphado tem parágrafos de verdade", _pont_para["paragrafos"] >= 8)
verificar("fixture paragraphado tem pontuação nativa", _pont_para["sinais_por_100_palavras"] > 5)
verificar("em arquivo paragraphado a 'linha mais longa' NÃO é o corpo",
          _int_para["cobertura_corpo_por_linha"] < 0.8)
verificar("e o perfil acusa isso como quebra, não em silêncio",
          any(not q["compativel"] for q in _int_para["quebras"]))
verificar("rótulos de fala nativos são detectados", bool(_int_para["rotulos_nativos"]))
verificar("ausência do marcador 'Transcrição Automática' é detectada",
          _int_para["marcador_cabecalho"] is False)

_int_yt = PS.eixo_integracao(PS.ler(FIXTURE))
verificar("a cobertura da maior linha discrimina os dois formatos",
          _int_yt["cobertura_corpo_por_linha"] > 3 * _int_para["cobertura_corpo_por_linha"] > 0.3)
# a REGRA em si, medida num corpo de linha única — no fixture pequeno o cabeçalho pesa e a
# cobertura fica abaixo do corte, o que é comportamento correto, não defeito do instrumento
_corpo = "olá eu sou o apresentador e hoje nós vamos falar sobre brahma. " * 40
_int_mono = PS.eixo_integracao({"texto": _corpo, "linhas": [_corpo]})
verificar("corpo de linha única é reconhecido como compatível",
          _int_mono["cobertura_corpo_por_linha"] >= 0.8 and
          all(q["compativel"] for q in _int_mono["quebras"] if "linha mais longa" in q["suposicao"]))

verificar("estágio do bruto de referência é 'bruto'",
          PS.estagio(REFERENCIA / "00-fonte" / "transcricao-bruta.txt")[0] == "bruto")
verificar("estágio de arquivo em upload/ é 'solto' (leitura de bruto)",
          PS.estagio(RAIZ / "upload" / "opcao-b.txt")[0] == "solto")

_regua = PS.superficies_kb(RAIZ / "KB-RC")
verificar("a régua carrega canônicos da KB", len(_regua["canonicos"]) > 500)
verificar("corrupção não pode ser também canônico (senão o eixo 3 mente)",
          not (set(_regua["variantes"]) & set(_regua["canonicos"])))
verificar("campo 'Variações' (equivalência conceitual) não entra como corrupção",
          L11.norm("humano") not in _regua["variantes"])
verificar("externos.csv é lido apesar dos comentários antes do cabeçalho",
          len(_regua["externos"]) > 0 and len(_regua["externos_canonicos"]) > 0)

_term_para = PS.eixo_terminologia(_para["texto"], _regua)
_term_yt = PS.eixo_terminologia(PS.ler(FIXTURE)["texto"], _regua)
verificar("os dois fixtures expõem as mesmas corrupções ao eixo 3",
          _term_para["variantes_stt_distintas"] >= 5 and _term_yt["variantes_stt_distintas"] >= 5)
verificar("livro-razão entra na régua quando o arquivo pertence a uma transcrição",
          len(PS.regua_para(_regua, REFERENCIA / "00-fonte" / "transcricao-bruta.txt",
                            RAIZ / "KB-RC")["proibidas"]) > len(_regua["proibidas"]))

# 13. o eixo de proveniência e a separação de ruído nos ausentes — o que o experimento de 16/09
#     descobriu na prática: os dois arquivos eram o MESMO reconhecimento de fala, e a lista de
#     "ausentes da base" de um STT pontuado enche de inicial de frase. Sem estes dois testes o
#     instrumento voltaria a confundir camada de reescrita com motor, e ruído com termo novo.
_linhas_aus = [{"forma": "Antigo Testamento", "ocorrencias": "1"},
               {"forma": "Uhum", "ocorrencias": "2"},
               {"forma": "Infelizmente", "ocorrencias": "1"},
               {"forma": "Kaagen", "ocorrencias": "2"}]
_txt_aus = L11.norm("antigo testamento Infelizmente infelizmente uhum uhum kaagen kaagen")
_cls = PS.classificar_ausentes(_linhas_aus, _txt_aus, PS.carregar_guarda())
verificar("janela de 2+ palavras vai para o balde n_grama", _cls["n_grama"]["formas"] == 1)
verificar("palavra do vocabulário-guarda não vira candidata a termo", _cls["comum"]["formas"] == 1)
verificar("inicial de frase é detectada como artefato de capitalização",
          _cls["capital"]["formas"] == 1 and _cls["capital"]["amostra"] == ["Infelizmente"])
verificar("o que sobra é candidata real", _cls["real"]["formas"] == 1
          and _cls["real"]["amostra"] == ["Kaagen"])

verificar("artigo grudado pela janela não conta como grafia diferente",
          PS.nucleo_grafia("do Paraíso Original") == "paraiso original")
verificar("grafias diferentes da mesma entidade continuam separadas",
          PS.nucleo_grafia("a constituicao centenaria") != PS.nucleo_grafia("Constituição Setenária"))

_pv_idem = PS.eixo_proveniencia(_perfis_fix[0], _perfis_fix[0])
verificar("arquivo comparado consigo mesmo tem divergência zero", _pv_idem["divergencia"] == 0.0)
verificar("e o veredito reconhece como a mesma fala", "MESMO reconhecimento" in _pv_idem["veredito"])
_pv_diff = PS.eixo_proveniencia(_perfis_fix[0], _perfis_fix[1])
verificar("textos diferentes não são declarados o mesmo ASR",
          _pv_diff["divergencia"] > 0.05 or "independentes" in _pv_diff["veredito"])

_seg_para = PS.eixo_pontuacao(_para["texto"], PS.paragrafos(_para["linhas"]), _para["linhas"])
verificar("segmento por quebra de linha conta turno de fala, não só parágrafo em branco",
          _seg_para["segmentos_por_quebra"] >= 8 and _seg_para["paragrafos"] < _seg_para["segmentos_por_quebra"])
verificar("censura por asterisco é detectada (a camada de reescrita mascara palavra)",
          PS.eixo_proveniencia(_perfis_fix[0], _perfis_fix[1])["mascarados_b"] == 1
          and PS.eixo_proveniencia(_perfis_fix[0], _perfis_fix[1])["mascarados_a"] == 0)
verificar("'bruto' e 'solto' são a mesma leitura (não dispara falso alarme de estágio)",
          PS.classe_de_leitura("bruto") == PS.classe_de_leitura("solto") != PS.classe_de_leitura("curado"))

# 14. arquitetura de três camadas: o corpo tem critério único (rc_leitura) e o portão G9 fiscaliza o
#     derivado. Medido sobre o caso real: a camada de reescrita mascarou "merda" ×3 e "bandido" ×2, e
#     o G1 continuaria verde porque confere sha256, não conteúdo. Um portão que não falha quando deve
#     é decoração, então cada modo de estragar tem seu teste.
import rc_leitura as RL  # noqa: E402
import rc_qa as QA14  # noqa: E402

REF = REFERENCIA / "00-fonte"
_mono = RL.localizar_corpo(RL.ler_texto(REF / "transcricao-bruta.txt"))
verificar("bruto do YouTube: corpo pelo marcador", _mono["criterio_corpo"] == "marcador")
verificar("e o número de palavras é o que já estava em metadados.yaml (18806)",
          _mono["corpo_palavras"] == 18806)
_pont = RL.localizar_corpo(RL.ler_texto(REF / "transcricao-pontuada.txt"))
verificar("derivado pontuado: corpo pelo marcador, não pela linha mais longa",
          _pont["criterio_corpo"] == "marcador" and _pont["corpo_segmentos"] > 280)
verificar("o critério antigo teria medido 771 palavras num corpo de 18 mil",
          _pont["corpo_palavras"] > 18000 and _pont["corpo_palavras"] > 20 * 771)
_sem_marca = RL.localizar_corpo("\n".join(["linha %d" % i for i in range(40)]))
verificar("arquivo paragraphado sem marcador cai no terceiro critério E avisa",
          _sem_marca["criterio_corpo"] == "arquivo-inteiro" and bool(_sem_marca["aviso_corpo"]))
_linha_unica = "palavra " * 500
verificar("linha única sem marcador usa o segundo critério, sem aviso",
          RL.localizar_corpo(_linha_unica)["criterio_corpo"] == "linha-mais-longa"
          and RL.localizar_corpo(_linha_unica)["aviso_corpo"] is None)

# pasta sintética mínima para exercitar o G9 nos dois sentidos
def _montar(tmp: Path, *, restaurada: bool = True, mascara_no_bloco: bool = False,
            sha: str = "", teto: str = "0.05", registrar: bool = True,
            derivado_extra: str = "") -> Path:
    # recheio: o teto de divergência é proporcional, e num texto de 30 palavras quatro máscaras
    # seriam 25% — o caso real tem 19 mil palavras, então o fixture também precisa ser longo
    recheio = " ".join(f"discurso{i}" for i in range(300))
    bruto = ("Título\nCanal\n\nTranscrição Automática\n"
             + recheio + " então o europeu via jesus como um bandido e o outro falou "
             "merda merda merda e depois seguiu o discurso normalmente até o fim da "
             "transmissão ao vivo\n")
    deriv = ("Título\nCanal\n\nTranscrição Automática\n"
             + recheio + ". Então, o europeu via Jesus como um b******. E o outro falou "
             "m****, m****, m****. E depois seguiu o discurso normalmente até o fim da "
             "transmissão ao vivo." + derivado_extra + "\n")
    (tmp / "00-fonte").mkdir(parents=True, exist_ok=True)
    (tmp / "20-blocos").mkdir(parents=True, exist_ok=True)
    (tmp / "00-fonte" / "transcricao-bruta.txt").write_text(bruto, encoding="utf-8")
    (tmp / "00-fonte" / "transcricao-pontuada.txt").write_text(deriv, encoding="utf-8")
    real = hashlib.sha256((tmp / "00-fonte" / "transcricao-pontuada.txt").read_bytes()).hexdigest()
    lista = ('    - mascara: "b******"\n      restaurar_para: bandido\n'
             '    - mascara: "m****"\n      restaurar_para: merda\n') if registrar else ""
    (tmp / "00-fonte" / "metadados.yaml").write_text(
        "slug: teste\nderivado:\n  arquivo: transcricao-pontuada.txt\n"
        f"  sha256: {sha or real}\n  divergencia_maxima: {teto}\n"
        f"  palavras_mascaradas:\n{lista}", encoding="utf-8")
    corpo = ("# Bloco 01\n\n**[GURU]** Então, o europeu via Jesus como um "
             + ("b******" if mascara_no_bloco else "bandido")
             + ". E o outro falou " + ("m****" if not restaurada and mascara_no_bloco else
                                      ("nada" if not restaurada else "merda, merda, merda"))
             + ".\n")
    (tmp / "20-blocos" / "bloco-01.md").write_text(corpo, encoding="utf-8")
    return tmp

with tempfile.TemporaryDirectory(prefix="rc-g9-") as _t:
    _p = _montar(Path(_t))
    _m = QA14.ler_metadados(_p)
    _st, _dt = QA14.g9_derivado_divergencia(_p, _m)
    verificar("G9 ok no caso bem formado", _st == QA14.OK)
    verificar("G9 nomeia as palavras que a camada mascarou",
              "bandido" in _dt and "merda" in _dt)
    _st2, _dt2 = QA14.g9_derivado_divergencia(_p, {**_m, "derivado": {**_m["derivado"],
                                                                      "sha256": "0" * 64}})
    verificar("G9 falha se o derivado foi trocado (sha256)", _st2 == QA14.FALHA
              and "ALTERADO" in _dt2)
    _st3, _dt3 = QA14.g9_derivado_divergencia(_p, {**_m, "derivado": {**_m["derivado"],
                                                                      "divergencia_maxima": "0.001"}})
    verificar("G9 falha se a divergência estoura o teto", _st3 == QA14.FALHA
              and "diverge" in _dt3)
with tempfile.TemporaryDirectory(prefix="rc-g9-") as _t:
    _p = _montar(Path(_t), restaurada=False)
    _st, _dt = QA14.g9_derivado_divergencia(_p, QA14.ler_metadados(_p))
    verificar("G9 falha se a palavra censurada não foi restaurada", _st == QA14.FALHA
              and "não foi restaurada" in _dt)
with tempfile.TemporaryDirectory(prefix="rc-g9-") as _t:
    _p = _montar(Path(_t), mascara_no_bloco=True)
    _st, _dt = QA14.g9_derivado_divergencia(_p, QA14.ler_metadados(_p))
    verificar("G9 falha se o asterisco chegou ao texto revisado", _st == QA14.FALHA
              and "copiada para o texto revisado" in _dt)
with tempfile.TemporaryDirectory(prefix="rc-g9-") as _t:
    _p = _montar(Path(_t), registrar=False)
    _st, _dt = QA14.g9_derivado_divergencia(_p, QA14.ler_metadados(_p))
    verificar("G9 falha se a máscara não está registrada em metadados.yaml", _st == QA14.FALHA
              and "não registrada" in _dt)
with tempfile.TemporaryDirectory(prefix="rc-g9-") as _t:
    _p = _montar(Path(_t), derivado_extra=" " + " ".join(f"enchimento{i}" for i in range(4000)))
    _st, _dt = QA14.g9_derivado_divergencia(_p, QA14.ler_metadados(_p))
    verificar("G9 falha se o derivado não é o mesmo texto (acréscimo massivo)", _st == QA14.FALHA)
verificar("G9 é N/A quando não há camada derivada",
          QA14.g9_derivado_divergencia(Path("/tmp"), {})[0] == QA14.NA)
_dv = QA14.divergencia_lexical("a b c d", "a b c d")
verificar("textos idênticos têm divergência zero", _dv["divergencia"] == 0.0)
_loc = QA14.localizar_mascaras("fala merda agora", "fala m**** agora", Counter({"merda": 1}))
verificar("a máscara é lida no bruto pelo contexto", _loc[0]["palavra_no_bruto"] == "merda")
_loc2 = QA14.localizar_mascaras("a b c d e f", "x y z m**** w q r", Counter())
verificar("sem alinhamento de contexto o portão não chuta palavra",
          _loc2[0]["palavra_no_bruto"] is None)

# o tempdir da seção 10 era criado e nunca removido — ficava um /tmp/rc-curadoria-* por execução
shutil.rmtree(_tmp, ignore_errors=True)
verificar("teste não deixa tempdir para trás", not _tmp.exists())

# --------------------------------------------------------------------------------------
print(f"\n{'=' * 66}\n{_ok} verificações ok, {len(_falhas)} falhas")
for f in _falhas:
    print("  -", f)
sys.exit(1 if _falhas else 0)
