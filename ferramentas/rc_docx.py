# -*- coding: utf-8 -*-
"""
rc_docx — montagem do documento revisado final (Passo 3 do Guia) de forma determinística.

Por que um script e não geração manual: o Guia exige tipografia fixa (parágrafos
justificados, entrelinha 1,5, corpo 12 pt), negrito na PRIMEIRA menção de cada termo
canônico, notas entre colchetes e expurgo de artefatos. Feito à mão, isso é caro e
inconsistente; feito por script, é reproduzível e auditável.

Entrada: um ou mais arquivos Markdown com os blocos revisados, na ordem, usando:
    # Título da transcrição
    ## Seção (opcional)
    parágrafo simples (linha em branco separa parágrafos)
    **Termo**            -> negrito explícito
    [NOTA: ...]          -> nota editorial (sai em itálico)

Com --lexico, o negrito da primeira menção de cada termo canônico é aplicado
automaticamente (e o negrito manual passa a ser opcional).

Uso:
    python ferramentas/rc_docx.py transcricoes/<slug>/20-blocos/bloco-*.md \
        --lexico transcricoes/<slug>/10-diagnostico/dossie-bloco.txt \
        --saida transcricoes/<slug>/30-produto/transcricao-revisada.docx \
        --titulo "Revelações Cósmicas Urgente — Jan Val Ellam" \
        --subtitulo "Transcrição revisada — Padronização terminológica conforme a Revelação Cósmica de Jan Val Ellam" \
        --validar transcricoes/<slug>/10-diagnostico/variantes-propostas.csv
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import rc_lexicon as L  # noqa: E402

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Pt, RGBColor
except ImportError as exc:  # pragma: no cover
    raise SystemExit("python-docx não instalado. Rode: pip install -r ferramentas/requirements.txt") from exc

# Uma nota pode CITAR a notação da KB, e essa notação tem colchetes dentro — ex.:
# RC-636 registra «o jogo de pósitrons [STT 'positelétron']». Sem tolerar um nível de
# aninhamento, o marcador fechava cedo demais e o rabo da nota voltava a ser corpo:
# no .docx saía sem itálico, e no G3 uma forma proibida citada como evidência depois do
# colchete interno daria falso positivo. Evidência do defeito: vídeo 2, bloco 3
# (`transcricoes/2026-09-12-alienigenas-e-humanos-entre-nos`), 17/09/2026.
CORPO_MARCADOR = r"(?:[^\[\]]|\[[^\]]*\])*"
NOTA_RE = re.compile(r"\[NOTA:" + CORPO_MARCADOR + r"\]")
# Marcadores editoriais que CITAM o bruto de propósito. O QA de sobrevivência de
# variantes não pode punir uma [NOTA] que documenta a forma ouvida ("Xavé" -> Javé),
# senão o revisor é incentivado a apagar a evidência em vez de registrá-la.
MARCADOR_RE = re.compile(r"\[(?:NOTA|A CONFIRMAR|INAUDÍVEL|ANÚNCIO):?" + CORPO_MARCADOR + r"\]")
NEGRITO_RE = re.compile(r"\*\*(.+?)\*\*")
# Ordem importa: a NOTA é casada antes dos asteriscos para que o markup interno
# (*título de livro*) não vire run separado — dentro da nota tudo já sai em itálico.
INLINE_RE = re.compile(NOTA_RE.pattern + r"|\*\*(.+?)\*\*|\*(.+?)\*")
ASTERISCO_RE = re.compile(r"\*+")


# --------------------------------------------------------------------------------------
# Preparo do texto
# --------------------------------------------------------------------------------------

def carregar_lexico(caminho: Path) -> list[str]:
    """Lê o dossê (saída de rc_diagnostico) e devolve as formas canônicas, da mais longa
    para a mais curta — para que 'Choque de Realidade' tenha prioridade sobre 'Choque'."""
    formas = set()
    for linha in caminho.read_text(encoding="utf-8").splitlines():
        partes = linha.split("|")
        if len(partes) < 2 or not partes[1].strip():
            continue
        nucleo = partes[1].strip()
        for p in re.split(r"\s*/\s*", nucleo):
            p = p.strip(" .")
            if len(p) > 3 and not p.isupper() or (p.isupper() and len(p) > 2):
                formas.add(p)
    return sorted(formas, key=lambda f: (-len(f), f))


def aplicar_negrito(texto: str, formas: list[str], ja_vistas: set[str]) -> str:
    """Marca com ** a primeira ocorrência de cada termo canônico."""
    for forma in formas:
        if forma.lower() in ja_vistas:
            continue
        padrao = re.compile(r"(?<![\w*])(" + re.escape(forma) + r")(?![\w*])", re.IGNORECASE)
        novo, n = padrao.subn(r"**\1**", texto, count=1)
        if n:
            texto = novo
            ja_vistas.add(forma.lower())
    return texto


def analisar_blocos(arquivos: list[Path], formas: list[str]) -> list[dict]:
    """Converte os Markdown de blocos revisados numa lista de itens estruturados."""
    itens: list[dict] = []
    vistas: set[str] = set()
    for arq in arquivos:
        for linha in arq.read_text(encoding="utf-8").splitlines():
            linha = linha.rstrip()
            if not linha.strip():
                continue
            if linha.startswith("### "):
                itens.append(dict(tipo="h3", texto=linha[4:].strip()))
            elif linha.startswith("## "):
                itens.append(dict(tipo="h2", texto=linha[3:].strip()))
            elif linha.startswith("# "):
                itens.append(dict(tipo="h1", texto=linha[2:].strip()))
            elif linha.startswith("> "):
                # linhas consecutivas de citação formam UM parágrafo
                if itens and itens[-1]["tipo"] == "citacao":
                    itens[-1]["texto"] += " " + linha[2:].strip()
                else:
                    itens.append(dict(tipo="citacao", texto=linha[2:].strip()))
            else:
                texto = aplicar_negrito(linha, formas, vistas) if formas else linha
                itens.append(dict(tipo="p", texto=texto))
    return itens


# --------------------------------------------------------------------------------------
# Renderização
# --------------------------------------------------------------------------------------

def _run(par, texto: str, fonte: str, tamanho: int) -> None:
    """Escreve um trecho aplicando **negrito**, *itálico* e [NOTA: ...] em itálico menor.

    Dentro de uma NOTA os asteriscos são descartados: a nota inteira já sai em itálico,
    então manter o markup só sujaria o produto de leitura com "*The Singularity Is Near*".
    """
    pos = 0
    for m in INLINE_RE.finditer(texto):
        if m.start() > pos:
            par.add_run(texto[pos:m.start()])
        if m.group(0).startswith("[NOTA"):
            r = par.add_run(ASTERISCO_RE.sub("", m.group(0)))
            r.italic = True
            r.font.size = Pt(tamanho - 1)
        elif m.group(1):
            r = par.add_run(m.group(1))
            r.bold = True
        else:
            r = par.add_run(m.group(2))
            r.italic = True
        pos = m.end()
    if pos < len(texto):
        par.add_run(texto[pos:])
    for r in par.runs:
        r.font.name = fonte
        r.font.size = r.font.size or Pt(tamanho)


def montar(itens: list[dict], saida: Path, titulo: str, subtitulo: str,
           fonte: str = "Calibri", tamanho: int = 12) -> Path:
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = fonte
    estilo.font.size = Pt(tamanho)
    pf = estilo.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(titulo)
    r.bold = True
    r.font.size = Pt(tamanho + 4)
    r.font.name = fonte
    if subtitulo:
        s = doc.add_paragraph()
        s.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rs = s.add_run(subtitulo)
        rs.italic = True
        rs.font.size = Pt(tamanho - 1)
        rs.font.name = fonte
    doc.add_paragraph()

    for item in itens:
        tipo, texto = item["tipo"], item["texto"]
        if tipo in ("h1", "h2", "h3"):
            nivel = int(tipo[1])
            # estilo Heading dá painel de navegação num documento de ~20 mil palavras;
            # a formatação do run é sobrescrita logo abaixo para manter a tipografia fixa
            # exigida pelo Guia (corpo 12 pt, preto, sem o azul padrão dos templates).
            p = doc.add_paragraph(style=f"Heading {nivel}")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(12 if nivel == 1 else 10)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            r = p.add_run(texto)
            r.bold = nivel <= 2
            r.italic = nivel == 3
            r.font.size = Pt(tamanho + (2 if nivel == 1 else 1 if nivel == 2 else 0))
            r.font.name = fonte
            r.font.color.rgb = RGBColor(0, 0, 0)
        elif tipo == "citacao":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(24)
            _run(p, texto, fonte, tamanho)
        else:
            p = doc.add_paragraph()
            _run(p, texto, fonte, tamanho)

    saida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(saida)
    return saida


# --------------------------------------------------------------------------------------
# Controle de qualidade
# --------------------------------------------------------------------------------------

# Só estas decisões obrigam o texto final. "recusada" e "aceita-parcial" existem
# justamente para registrar que a variante sobrevive POR DECISÃO HUMANA — cobrar a
# substituição nesses casos empurraria o revisor a corromper o texto para agradar o QA.
COBRADAS = {"aceita"}


def validar(itens: list[dict], csv_variantes: Path | None) -> list[str]:
    """QA: procura no texto final variantes cuja substituição foi ADJUDICADA como aceita.

    Se o CSV traz a coluna `adjudicacao` (gerada pelo registro de decisões da revisão),
    apenas as linhas `aceita` são cobradas. Sem a coluna, vale o comportamento antigo
    (`status_aprovacao` em aprovada/proposta), para não quebrar pipelines legados.
    """
    texto = " ".join(i["texto"] for i in itens)
    # expurga o conteúdo dos marcadores editoriais: eles citam o bruto de propósito
    nb = L.norm(MARCADOR_RE.sub(" ", texto))
    problemas = []
    if not csv_variantes or not csv_variantes.exists():
        return problemas
    with csv_variantes.open(encoding="utf-8-sig", newline="") as fh:
        for linha in csv.DictReader(fh):
            if "adjudicacao" in linha:
                if (linha.get("adjudicacao") or "").strip().lower() not in COBRADAS:
                    continue
            else:
                if linha.get("status_aprovacao", "").lower() not in {"aprovada", "proposta"}:
                    continue
                if linha.get("classe") not in ("variante", "truncamento", "semente-guia"):
                    continue
            v = L.norm(linha["variante"])
            if not v:
                continue
            achados = len(re.findall(r"\b" + re.escape(v) + r"\b", nb))
            if achados:
                problemas.append(f"{achados}x '{linha['variante']}' -> deveria ser "
                                 f"'{linha['canonico_proposto']}' ({linha.get('codigo_base','')})")
    return problemas


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Monta o DOCX revisado a partir dos blocos Markdown.")
    ap.add_argument("blocos", nargs="+", type=Path)
    ap.add_argument("--saida", type=Path, required=True)
    ap.add_argument("--titulo", required=True)
    ap.add_argument("--subtitulo", default="")
    ap.add_argument("--lexico", type=Path, default=None,
                    help="dossie-bloco.txt do diagnóstico (ativa negrito automático de 1ª menção)")
    ap.add_argument("--fonte", default="Calibri")
    ap.add_argument("--tamanho", type=int, default=12)
    ap.add_argument("--validar", type=Path, default=None,
                    help="variantes-propostas.csv para QA de sobrevivência de variantes")
    args = ap.parse_args(argv)

    formas = carregar_lexico(args.lexico) if args.lexico else []
    itens = analisar_blocos(args.blocos, formas)
    caminho = montar(itens, args.saida, args.titulo, args.subtitulo, args.fonte, args.tamanho)
    palavras = sum(len(i["texto"].split()) for i in itens)
    notas = sum(len(NOTA_RE.findall(i["texto"])) for i in itens)
    negritos = sum(i["texto"].count("**") // 2 for i in itens)
    print(f"[ok] {caminho}")
    print(f"     blocos={len(args.blocos)} itens={len(itens)} palavras={palavras} "
          f"negritos={negritos} notas={notas} formas_no_lexico={len(formas)}")
    problemas = validar(itens, args.validar)
    if problemas:
        print(f"[qa] {len(problemas)} variantes ADJUDICADAS como aceitas ainda sobrevivem no texto:")
        for p in problemas[:20]:
            print("     -", p)
        return 1
    print("[qa] OK: nenhuma variante adjudicada como 'aceita' sobrevive no texto final."
          " (recusas e aceitações parciais estão registradas na coluna adjudicacao)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
