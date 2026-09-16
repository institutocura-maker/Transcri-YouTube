# -*- coding: utf-8 -*-
"""
md_para_docx — converte um documento Markdown institucional em .docx legível.

Motivo: o contrato editorial do projeto exige DOIS produtos para cada entrega —
o .md (versionamento, inegociável) e o .docx (produto de leitura). O `rc_docx.py`
cuida da transcrição revisada (tipografia fixa, negrito de primeira menção); este
script cuida dos documentos de governança: Guia, pareceres, relatórios de conflito.

Suporta: títulos #/##/###/####, parágrafos, **negrito**, *itálico*, `código`,
tabelas | a | b |, listas - e 1., citações >, linhas ---, e blocos ``` código ```.

Uso:
    python ferramentas/md_para_docx.py analise/RESOLUCAO-DE-CONFLITOS.md \
        --saida analise/RESOLUCAO-DE-CONFLITOS.docx \
        --titulo "Resolução de Conflitos da Base Terminológica" \
        --subtitulo "Guia de Revisão v2 — Anexo I"
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.exit("dependência ausente: pip install python-docx")

AZUL = RGBColor(0x1F, 0x3B, 0x57)
CINZA = RGBColor(0x55, 0x55, 0x55)

INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\n]+\*|`[^`]+`)")


def _escrever_inline(par, texto: str) -> None:
    """Aplica negrito/itálico/código dentro de um parágrafo."""
    for pedaco in INLINE.split(texto):
        if not pedaco:
            continue
        if pedaco.startswith("**") and pedaco.endswith("**"):
            run = par.add_run(pedaco[2:-2])
            run.bold = True
        elif pedaco.startswith("__") and pedaco.endswith("__"):
            run = par.add_run(pedaco[2:-2])
            run.bold = True
        elif pedaco.startswith("*") and pedaco.endswith("*") and len(pedaco) > 2:
            run = par.add_run(pedaco[1:-1])
            run.italic = True
        elif pedaco.startswith("`") and pedaco.endswith("`"):
            run = par.add_run(pedaco[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            par.add_run(pedaco)


def _celula(texto: str) -> str:
    return re.sub(r"\*\*|\*|`", "", texto).strip()


def converter(md: Path, docx_saida: Path, titulo: str = "", subtitulo: str = "",
              fonte: str = "Calibri", tamanho: int = 11) -> None:
    linhas = md.read_text(encoding="utf-8").splitlines()
    doc = Document()

    estilo = doc.styles["Normal"]
    estilo.font.name = fonte
    estilo.font.size = Pt(tamanho)
    estilo.paragraph_format.space_after = Pt(6)
    estilo.paragraph_format.line_spacing = 1.15

    if titulo:
        p = doc.add_paragraph()
        run = p.add_run(titulo)
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = AZUL
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitulo:
        p = doc.add_paragraph()
        run = p.add_run(subtitulo)
        run.italic = True
        run.font.size = Pt(11)
        run.font.color.rgb = CINZA
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if titulo or subtitulo:
        doc.add_paragraph()

    i = 0
    while i < len(linhas):
        linha = linhas[i].rstrip()

        # bloco de código
        if linha.startswith("```"):
            i += 1
            buf = []
            while i < len(linhas) and not linhas[i].startswith("```"):
                buf.append(linhas[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(buf))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            p.paragraph_format.left_indent = Pt(18)
            continue

        # tabela
        if linha.startswith("|") and i + 1 < len(linhas) and re.match(r"^\|[\s:\-|]+\|$", linhas[i + 1].strip()):
            cabecalho = [c for c in linha.strip().strip("|").split("|")]
            i += 2
            corpo = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                corpo.append([c for c in linhas[i].strip().strip("|").split("|")])
                i += 1
            ncol = len(cabecalho)
            tabela = doc.add_table(rows=1, cols=ncol)
            tabela.style = "Light Grid Accent 1"
            tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, txt in enumerate(cabecalho):
                cel = tabela.rows[0].cells[j]
                cel.text = ""
                run = cel.paragraphs[0].add_run(_celula(txt))
                run.bold = True
                run.font.size = Pt(tamanho - 1)
            for linha_tab in corpo:
                cells = tabela.add_row().cells
                for j in range(ncol):
                    txt = linha_tab[j] if j < len(linha_tab) else ""
                    cells[j].text = ""
                    run = cells[j].paragraphs[0].add_run(_celula(txt))
                    run.font.size = Pt(tamanho - 1)
            doc.add_paragraph()
            continue

        vazio = not linha.strip()
        if vazio:
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", linha)
        if m:
            nivel = len(m.group(1))
            texto = m.group(2).strip()
            if nivel == 1:
                p = doc.add_paragraph()
                run = p.add_run(re.sub(r"\*\*", "", texto))
                run.bold = True
                run.font.size = Pt(16)
                run.font.color.rgb = AZUL
                p.paragraph_format.space_before = Pt(14)
            else:
                p = doc.add_paragraph()
                run = p.add_run(re.sub(r"\*\*|\*", "", texto))
                run.bold = True
                run.font.size = Pt(14 - nivel)
                run.font.color.rgb = AZUL if nivel <= 3 else CINZA
                p.paragraph_format.space_before = Pt(10)
            i += 1
            continue

        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", linha.strip()):
            i += 1
            continue

        m = re.match(r"^>\s?(.*)$", linha)
        if m:
            p = doc.add_paragraph()
            _escrever_inline(p, m.group(1))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = CINZA
            p.paragraph_format.left_indent = Pt(24)
            i += 1
            continue

        m = re.match(r"^(\s*)[-*+]\s+(.*)$", linha)
        if m:
            recuo = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet")
            _escrever_inline(p, m.group(2))
            p.paragraph_format.left_indent = Pt(18 + 18 * recuo)
            i += 1
            continue

        m = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", linha)
        if m:
            p = doc.add_paragraph(style="List Number")
            _escrever_inline(p, m.group(3))
            i += 1
            continue

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _escrever_inline(p, linha)
        i += 1

    docx_saida.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_saida))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Converte Markdown institucional em .docx.")
    ap.add_argument("markdown", type=Path)
    ap.add_argument("--saida", type=Path, required=True)
    ap.add_argument("--titulo", default="")
    ap.add_argument("--subtitulo", default="")
    ap.add_argument("--fonte", default="Calibri")
    ap.add_argument("--tamanho", type=int, default=11)
    args = ap.parse_args(argv)

    md = args.markdown if args.markdown.is_absolute() else Path.cwd() / args.markdown
    if not md.exists():
        md = Path(__file__).resolve().parent.parent / args.markdown
    converter(md, args.saida, titulo=args.titulo, subtitulo=args.subtitulo,
              fonte=args.fonte, tamanho=args.tamanho)
    print(f"[ok] {args.saida} gerado a partir de {md.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
