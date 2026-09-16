# -*- coding: utf-8 -*-
"""rc_qa — os oito portões de qualidade de uma transcrição.

Por que um script só para isto: o QA existia, mas espalhado e dependente de memória.
Cada portão responde a um modo concreto de o trabalho estragar — e todos já aconteceram
nesta ou noutra forma durante a primeira transcrição:

    G1  bruto intacto           alguém "só corrigiu um errinho" no arquivo sagrado
    G2  blocos íntegros         bloco sem título, sem falante, com comentário HTML
    G3  formas proibidas        canônico violado no texto corrido (a KB diz NUNCA "Sofia")
    G4  ledger fechado          linha de adjudicação sem decisão — o revisor não terminou
    G5  validador               variante adjudicada como aceita sobreviveu
    G6  produto reproduzível    o .docx publicado não corresponde aos .md
    G7  índice consistente      o catálogo diz uma coisa, a pasta diz outra
    G8  higiene                 arquivo pesado, nome com espaço, ~$trava do Word

Três estados por portão: OK, FALHA e N/A (o portão não se aplica ao estágio em que a
transcrição está — não faz sentido cobrar .docx de quem ainda não revisou nada).

Uso:
    python ferramentas/rc_qa.py transcricoes/<slug>
    python ferramentas/rc_qa.py --tudo
    python ferramentas/rc_qa.py transcricoes/<slug> --portao G3
    python ferramentas/rc_qa.py transcricoes/<slug> --json
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import rc_docx as DX  # noqa: E402
import rc_lexicon as L  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
OK, FALHA, NA = "OK", "FALHA", "N/A"
# Marcadores editoriais citam o bruto de propósito: o conteúdo deles nunca é erro.
MARCADOR_RE = DX.MARCADOR_RE
ROTULO_RE = re.compile(r"\*\*\[[A-ZÀ-Ý][A-ZÀ-Ý ?]*\]\*\*")
ACENTUADO_RE = re.compile(r"[^\x00-\x7F]")
LIMITE_PESO = 5 * 1024 * 1024  # 5 MB fora do LFS reprova (Plano §7)


# --------------------------------------------------------------------------------------
# Leitura de metadados (subconjunto YAML do modelo — sem dependência externa)
# --------------------------------------------------------------------------------------

def ler_metadados(pasta: Path) -> dict:
    """Lê o `metadados.yaml` do modelo: dois níveis de indentação, valores escalares.

    Não é um parser YAML e não pretende ser: instalar PyYAML para ler um arquivo que
    este projeto mesmo define é dependência sem ganho. Se o modelo ganhar listas ou
    aninhamento profundo, troque isto por PyYAML e apague esta função.
    """
    arq = pasta / "00-fonte" / "metadados.yaml"
    if not arq.exists():
        return {}
    raiz: dict = {}
    atual = raiz
    for linha in arq.read_text(encoding="utf-8").splitlines():
        if not linha.strip() or linha.lstrip().startswith("#"):
            continue
        recuo = len(linha) - len(linha.lstrip(" "))
        corpo = linha.split("#")[0].strip() if not linha.strip().startswith("#") else ""
        if not corpo or ":" not in corpo:
            continue
        chave, _, valor = corpo.partition(":")
        chave, valor = chave.strip(), valor.strip().strip('"')
        if recuo == 0:
            atual = raiz
            if valor == "":
                raiz[chave] = {}
                atual = raiz[chave]
            else:
                raiz[chave] = valor
        elif isinstance(atual, dict):
            atual[chave] = valor
    return raiz


def _inteiro(v, padrao=0) -> int:
    try:
        return int(str(v).replace(".", "").strip())
    except (TypeError, ValueError):
        return padrao


def estagio(pasta: Path) -> str:
    """Estágio em que a transcrição está, pelos artefatos que existem de fato.

    Status aspiracional é a forma mais comum de autoengano em acervo: aqui quem diz
    o estágio é o disco, não o catálogo. O portão G7 compara os dois.
    """
    if (pasta / "40-devolucao" / "devolucao-a-kb.md").exists():
        return "40-devolvida"
    if (pasta / "30-produto" / "transcricao-revisada.docx").exists():
        return "30-revisada"
    if list((pasta / "20-blocos").glob("bloco-*.md")) if (pasta / "20-blocos").exists() else []:
        return "20-em-revisao"
    if (pasta / "10-diagnostico" / "variantes-propostas.csv").exists():
        return "10-em-diagnostico"
    return "00-nova"


# --------------------------------------------------------------------------------------
# Insumos compartilhados
# --------------------------------------------------------------------------------------

def _blocos(pasta: Path) -> list[Path]:
    d = pasta / "20-blocos"
    return sorted(d.glob("bloco-*.md")) if d.exists() else []


def _texto_blocos(pasta: Path) -> str:
    return "\n".join(b.read_text(encoding="utf-8") for b in _blocos(pasta))


def _ledger(pasta: Path) -> list[dict]:
    arq = pasta / "10-diagnostico" / "variantes-propostas.csv"
    if not arq.exists():
        return []
    with arq.open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def formas_proibidas(pasta: Path, kb: Path) -> dict[str, str]:
    """Forma proibida -> motivo. Vem de duas fontes, ambas de verdade:

    1. a **Quarentena** das fichas da KB-RC (`NUNCA "Sofia"`), que é a fonte de verdade;
    2. o **livro-razão** desta transcrição: variante adjudicada como `aceita` não pode
       sobreviver no texto corrido.

    De propósito NÃO entra aqui a lista inteira de sementes STT: formas curtas como
    "Nick" colidem com palavra comum e só fazem sentido com contexto — o motor já cuida
    disso, e o revisor registra a decisão no ledger.
    """
    proibidas: dict[str, str] = {}
    try:
        import rc_kb as KB
        for ficha in KB.carregar_fichas(str(kb)).values():
            q = ficha.quarentena or ""
            for m in re.finditer(r'NUNCA\s+\**"([^"\n]+?)"\**', q):
                proibidas.setdefault(m.group(1).strip(), f'Quarentena de {ficha.codigo}: NUNCA "{m.group(1).strip()}"')
    except Exception as exc:  # noqa: BLE001 — KB indisponível não pode derrubar o QA inteiro
        print(f"[aviso] não foi possível ler a Quarentena da KB: {exc}", file=sys.stderr)
    for linha in _ledger(pasta):
        if (linha.get("adjudicacao") or "").strip().lower() == "aceita":
            v = linha["variante"].strip()
            if v and len(v) > 2:
                proibidas.setdefault(v, f"ledger: '{v}' adjudicada como aceita -> {linha.get('canonico_proposto','')}")
    return proibidas


# --------------------------------------------------------------------------------------
# Portões
# --------------------------------------------------------------------------------------

def g1_bruto_intacto(pasta: Path, meta: dict) -> tuple[str, str]:
    arq = pasta / "00-fonte" / "transcricao-bruta.txt"
    if not arq.exists():
        return FALHA, "00-fonte/transcricao-bruta.txt não existe"
    esperado = (meta.get("bruto", {}) or {}).get("sha256", "")
    real = hashlib.sha256(arq.read_bytes()).hexdigest()
    if not esperado:
        return FALHA, "metadados.yaml não registra o sha256 do bruto"
    if esperado.lower() != real:
        return FALHA, f"bruto ALTERADO: sha256 {real[:12]}… != {esperado[:12]}… registrado"
    return OK, f"sha256 {real[:12]}… confere ({arq.stat().st_size} bytes)"


def g2_blocos_integros(pasta: Path, meta: dict) -> tuple[str, str]:
    blocos = _blocos(pasta)
    if not blocos:
        return NA, "nenhum bloco em 20-blocos/ (transcrição ainda não revisada)"
    problemas = []
    for b in blocos:
        texto = b.read_text(encoding="utf-8")
        if not re.match(r"^bloco-\d{2,}\.md$", b.name):
            problemas.append(f"{b.name}: nome fora do padrão bloco-NN.md (ordenação quebra com >9 blocos)")
        if not re.search(r"^## Bloco \d+", texto, re.M):
            problemas.append(f"{b.name}: sem título '## Bloco N — tema'")
        if not ROTULO_RE.search(texto):
            problemas.append(f"{b.name}: sem nenhum rótulo de fala")
        if "<!--" in texto:
            problemas.append(f"{b.name}: comentário HTML (o montador não o conhece)")
        if re.search(r"\[RASCUNHO|-rascunho\.md$", texto + b.name, re.I):
            problemas.append(f"{b.name}: marcador de rascunho não removido")
    if problemas:
        return FALHA, "; ".join(problemas[:5]) + ("…" if len(problemas) > 5 else "")
    rotulos = len(ROTULO_RE.findall(_texto_blocos(pasta)))
    return OK, f"{len(blocos)} blocos, {rotulos} rótulos de fala, nomes padronizados"


def g3_formas_proibidas(pasta: Path, meta: dict, kb: Path) -> tuple[str, str]:
    blocos = _blocos(pasta)
    if not blocos:
        return NA, "nada a varrer"
    texto = MARCADOR_RE.sub(" ", _texto_blocos(pasta))
    proibidas = formas_proibidas(pasta, kb)
    if not proibidas:
        return NA, "nenhuma forma proibida derivável (KB sem Quarentena e ledger sem linhas aceitas)"
    achados = []
    for forma, motivo in proibidas.items():
        # fronteira de palavra é essencial: sem ela, "Demiurg" casa dentro de "Demiurgo"
        # e "enoteísmo" dentro de "henoteísmo", e o QA passa a acusar o texto correto.
        # L.fronteira (e não \b puro) porque forma que começa em não-palavra — "/Kaggen",
        # RC-948 — não tem \b inicial e passaria invisível pelo portão.
        n = len(re.findall(L.fronteira(forma), texto, re.I))
        if n:
            achados.append(f"{n}x '{forma}' ({motivo})")
    if achados:
        return FALHA, "; ".join(achados[:6]) + ("…" if len(achados) > 6 else "")
    return OK, f"{len(proibidas)} formas proibidas varridas com fronteira de palavra, 0 ocorrências"


def g4_ledger_fechado(pasta: Path, meta: dict) -> tuple[str, str]:
    linhas = _ledger(pasta)
    if not linhas:
        return NA, "10-diagnostico/variantes-propostas.csv ausente (diagnóstico não rodou)"
    if "adjudicacao" not in linhas[0]:
        return FALHA, "o ledger não tem a coluna adjudicacao — decisões não registradas"
    pendentes = [l["variante"] for l in linhas
                 if not (l.get("adjudicacao") or "").strip()
                 or (l.get("adjudicacao") or "").strip().upper() == "REVISAR"]
    if pendentes:
        return FALHA, f"{len(pendentes)} linhas sem decisão: {', '.join(pendentes[:5])}"
    conta: dict[str, int] = {}
    for l in linhas:
        k = l["adjudicacao"].strip().lower()
        conta[k] = conta.get(k, 0) + 1
    resumo = ", ".join(f"{v} {k}" for k, v in sorted(conta.items(), key=lambda x: -x[1]))
    return OK, f"{len(linhas)} linhas decididas ({resumo})"


def g5_validador(pasta: Path, meta: dict) -> tuple[str, str]:
    blocos = _blocos(pasta)
    csv_ledger = pasta / "10-diagnostico" / "variantes-propostas.csv"
    if not blocos or not csv_ledger.exists():
        return NA, "sem blocos ou sem ledger"
    lexico = pasta / "10-diagnostico" / "dossie-bloco.txt"
    formas = DX.carregar_lexico(lexico) if lexico.exists() else []
    itens = DX.analisar_blocos(blocos, formas)
    problemas = DX.validar(itens, csv_ledger)
    if problemas:
        return FALHA, f"{len(problemas)} variantes aceitas sobrevivem: " + "; ".join(problemas[:4])
    return OK, "nenhuma variante adjudicada como aceita sobrevive no texto"


def g6_produto_reproduzivel(pasta: Path, meta: dict) -> tuple[str, str]:
    blocos = _blocos(pasta)
    produto = pasta / "30-produto" / "transcricao-revisada.docx"
    if not blocos:
        return NA, "sem blocos"
    if not produto.exists():
        return NA, "30-produto/transcricao-revisada.docx ainda não foi montado"
    from docx import Document
    lexico = pasta / "10-diagnostico" / "dossie-bloco.txt"
    formas = DX.carregar_lexico(lexico) if lexico.exists() else []
    itens = DX.analisar_blocos(blocos, formas)
    titulo = (meta.get("titulo") or pasta.name)
    with tempfile.TemporaryDirectory() as tmp:
        novo = DX.montar(itens, Path(tmp) / "r.docx", titulo, "", "Calibri", 12)
        a = [p.text for p in Document(str(produto)).paragraphs]
        b = [p.text for p in Document(str(novo)).paragraphs]
    # O cabeçalho (título e subtítulo) é informado por argumento na montagem e pode
    # diferir legitimamente; a comparação vale do primeiro bloco em diante.
    inicio_a = next((i for i, t in enumerate(a) if t.startswith("Bloco ")), 0)
    inicio_b = next((i for i, t in enumerate(b) if t.startswith("Bloco ")), 0)
    a, b = a[inicio_a:], b[inicio_b:]
    if len(a) != len(b):
        return FALHA, f".docx versionado tem {len(a)} parágrafos de corpo; regenerado dos .md dá {len(b)}"
    diffs = [i for i, (x, y) in enumerate(zip(a, b)) if x != y]
    if diffs:
        return FALHA, f"{len(diffs)} parágrafos diferem (primeiro: #{diffs[0]}) — o produto está fora de sincronia com os blocos"
    return OK, f"{len(a)} parágrafos de corpo idênticos aos regenerados a partir dos .md"


def g7_indice_consistente(pasta: Path, meta: dict) -> tuple[str, str]:
    indice = pasta.parent / "_indice.csv"
    if not indice.exists():
        return FALHA, "transcricoes/_indice.csv não existe"
    with indice.open(encoding="utf-8-sig", newline="") as fh:
        linhas = {l["slug"]: l for l in csv.DictReader(fh)}
    slug = pasta.name
    if slug not in linhas:
        return FALHA, f"slug '{slug}' não está no catálogo"
    reg = linhas[slug]
    disco = estagio(pasta)
    if reg.get("estatus") != disco:
        return FALHA, f"catálogo diz '{reg.get('estatus')}', o disco diz '{disco}'"
    problemas = []
    nblocos = len(_blocos(pasta))
    if _inteiro(reg.get("blocos")) != nblocos:
        problemas.append(f"catálogo declara {reg.get('blocos')} blocos, existem {nblocos}")
    if meta.get("slug") and meta["slug"] != slug:
        problemas.append(f"metadados.yaml declara slug '{meta['slug']}'")
    if (meta.get("bruto", {}) or {}).get("arquivo") not in (None, "", "transcricao-bruta.txt"):
        problemas.append("metadados.yaml aponta outro arquivo de bruto")
    if problemas:
        return FALHA, "; ".join(problemas)
    return OK, f"status '{disco}' confere com o disco e com o catálogo; {nblocos} blocos"


def g8_higiene(pasta: Path, meta: dict) -> tuple[str, str]:
    problemas = []
    alvos = [pasta, RAIZ / "ferramentas"]
    for alvo in alvos:
        for arq in alvo.rglob("*"):
            if not arq.is_file() or ".git" in arq.parts:
                continue
            nome = arq.relative_to(RAIZ).as_posix()
            if arq.stat().st_size > LIMITE_PESO:
                problemas.append(f"{nome}: {arq.stat().st_size // 1024 // 1024} MB (limite 5 MB fora do LFS)")
            if arq.name.startswith("~$"):
                problemas.append(f"{nome}: trava do Office não deve ser versionada")
            if " " in arq.name:
                problemas.append(f"{nome}: espaço no nome")
            elif ACENTUADO_RE.search(arq.name):
                problemas.append(f"{nome}: acento no nome (caminho de máquina é ASCII)")
    if problemas:
        return FALHA, "; ".join(problemas[:6]) + ("…" if len(problemas) > 6 else "")
    return OK, "nenhum arquivo acima de 5 MB, sem travas do Office, nomes ASCII e sem espaço"


PORTOES = [
    ("G1", "bruto intacto", g1_bruto_intacto),
    ("G2", "blocos íntegros", g2_blocos_integros),
    ("G3", "formas proibidas", g3_formas_proibidas),
    ("G4", "ledger fechado", g4_ledger_fechado),
    ("G5", "validador rc_docx", g5_validador),
    ("G6", "produto reproduzível", g6_produto_reproduzivel),
    ("G7", "índice consistente", g7_indice_consistente),
    ("G8", "higiene", g8_higiene),
]


def rodar(pasta: Path, kb: Path, so_portao: str | None = None) -> list[dict]:
    pasta = pasta.resolve()
    meta = ler_metadados(pasta)
    resultados = []
    for codigo, nome, funcao in PORTOES:
        if so_portao and codigo != so_portao:
            continue
        try:
            if codigo == "G3":
                status, detalhe = funcao(pasta, meta, kb)
            else:
                status, detalhe = funcao(pasta, meta)
        except Exception as exc:  # noqa: BLE001 — um portão quebrado não pode esconder os outros
            status, detalhe = FALHA, f"erro interno no portão: {exc}"
        resultados.append(dict(portao=codigo, nome=nome, status=status, detalhe=detalhe))
    return resultados


def imprimir(pasta: Path, resultados: list[dict]) -> None:
    print(f"\n{pasta.name}")
    for r in resultados:
        marca = {"OK": "ok ", "FALHA": "FALHA", "N/A": "n/a"}[r["status"]]
        print(f"  [{r['portao']}] {r['nome']:<22} {marca:<5} {r['detalhe']}")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Portões de qualidade de uma transcrição (Plano §8.2).")
    ap.add_argument("pasta", type=Path, nargs="?", help="transcricoes/<slug>")
    ap.add_argument("--tudo", action="store_true", help="varre transcricoes/*/")
    ap.add_argument("--portao", help="roda um só portão (G1…G8)")
    ap.add_argument("--kb", type=Path, default=RAIZ / "KB-RC")
    ap.add_argument("--json", action="store_true", help="saída máquina-legível para o CI")
    args = ap.parse_args(argv)

    if args.tudo:
        pastas = sorted(p for p in (RAIZ / "transcricoes").iterdir()
                        if p.is_dir() and not p.name.startswith("_"))
    elif args.pasta:
        pastas = [args.pasta if args.pasta.is_absolute() else RAIZ / args.pasta]
    else:
        ap.error("informe a pasta da transcrição ou use --tudo")

    saida, falhas = [], 0
    for pasta in pastas:
        resultados = rodar(pasta, args.kb, args.portao)
        saida.append(dict(slug=pasta.name, estagio=estagio(pasta), portoes=resultados))
        if any(r["status"] == FALHA for r in resultados):
            falhas += 1
        if not args.json:
            imprimir(pasta, resultados)

    if args.json:
        print(json.dumps(saida, ensure_ascii=False, indent=1))
    if not args.json:
        print()
    return 1 if falhas else 0


if __name__ == "__main__":
    raise SystemExit(main())
